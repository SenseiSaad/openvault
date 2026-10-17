"""Features 2 & 3 — AI workflow actions triggered by detected chat intents.

`handle_ai_workflow_intent(intent_data, user)` is the single service-layer entry
point called by the RAG stream once the assistant has emitted an intent block:

  * schedule_meeting -> create a Google Calendar event (real when credentials are
    configured, otherwise a graceful simulation) and send a confirmation email.
  * leave_request    -> generate an auto-filled .docx leave application with
    python-docx, store it under MEDIA_ROOT/forms/<user_id>/, and return a
    download URL served by an authenticated, ownership-checked view.

Design rules:
  * Every heavy/optional dependency (googleapiclient) is imported lazily and
    failures degrade gracefully — the chat must never break because a meeting
    could not be booked or an email could not be sent.
  * No secrets are hardcoded. Calendar credentials and SMTP settings come from
    Django settings / environment only.
  * The function NEVER raises; it always returns a JSON-serialisable dict (or
    None for a no-op intent) describing what happened.
"""
from __future__ import annotations

import logging
import os
import re
import uuid
from datetime import datetime, timedelta

from django.conf import settings
from django.core.mail import send_mail
from django.utils import timezone

log = logging.getLogger(__name__)

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._-]+")


def handle_ai_workflow_intent(intent_data: dict | None, user) -> dict | None:
    """Dispatch a parsed intent to the matching workflow. Returns a result dict
    describing the outcome, or None when there is nothing to do."""
    if not isinstance(intent_data, dict):
        return None
    intent = (intent_data.get("intent") or "none").strip()
    data = intent_data.get("data") or {}
    if not isinstance(data, dict):
        data = {}

    try:
        if intent == "schedule_meeting":
            return _schedule_meeting(data, user)
        if intent == "leave_request":
            return _leave_request(data, user)
    except Exception as exc:  # never let a workflow failure break the chat
        log.exception("workflow intent %s failed", intent)
        return {
            "action": intent,
            "status": "error",
            "message": "The assistant could not complete that action automatically.",
            "detail": str(exc)[:200],
        }
    return None


# ─────────────────────────── schedule_meeting ───────────────────────────────

def _schedule_meeting(data: dict, user) -> dict:
    title = (data.get("title") or "Meeting").strip()
    date = (data.get("date") or "").strip()
    time = (data.get("time") or "").strip()
    participants = data.get("participants") or []
    if not isinstance(participants, list):
        participants = [str(participants)]
    participants = [str(p).strip() for p in participants if str(p).strip()]

    start_dt, end_dt = _meeting_window(date, time)
    calendar = _create_calendar_event(title, start_dt, end_dt, participants, user)
    email = _send_meeting_email(title, start_dt, participants, user, calendar)

    when = start_dt.strftime("%d %b %Y, %H:%M") if start_dt else (date or "a time to be confirmed")
    if calendar["status"] == "created":
        message = f"Done. '{title}' is booked for {when}."
    else:
        message = f"Noted. I prepared '{title}' for {when} (calendar sync is not configured yet)."

    return {
        "action": "schedule_meeting",
        "status": "ok" if calendar["status"] == "created" else "simulated",
        "message": message,
        "title": title,
        "when": when,
        "participants": participants,
        "calendar": calendar,
        "email": email,
    }


def _meeting_window(date: str, time: str):
    """Parse date (YYYY-MM-DD) + time (HH:MM) into a (start, end) pair. End is one
    hour after start. Returns (None, None) if the date can't be parsed."""
    if not date:
        return None, None
    fmt = "%Y-%m-%d %H:%M" if time else "%Y-%m-%d"
    raw = f"{date} {time}".strip()
    try:
        start = datetime.strptime(raw, fmt)
    except ValueError:
        return None, None
    if timezone.is_naive(start) and settings.USE_TZ:
        start = timezone.make_aware(start, timezone.get_current_timezone())
    return start, start + timedelta(hours=1)


def _create_calendar_event(title, start_dt, end_dt, participants, user) -> dict:
    """Create a real Google Calendar event when possible; otherwise simulate.
    Always returns a dict with a `status` of 'created' | 'simulated' | 'error'."""
    creds_path = settings.GOOGLE_CALENDAR_CREDENTIALS
    if not start_dt:
        return {"status": "simulated", "reason": "no valid date/time supplied"}
    if not creds_path or not os.path.exists(creds_path):
        return {"status": "simulated", "reason": "calendar credentials not configured"}
    try:
        from google.oauth2 import service_account  # lazy: optional dependency
        from googleapiclient.discovery import build

        scopes = ["https://www.googleapis.com/auth/calendar"]
        credentials = service_account.Credentials.from_service_account_file(creds_path, scopes=scopes)
        service = build("calendar", "v3", credentials=credentials, cache_discovery=False)
        tz = settings.TIME_ZONE
        body = {
            "summary": title,
            "description": f"Scheduled via OpenVault assistant for {user.email}.",
            "start": {"dateTime": start_dt.isoformat(), "timeZone": tz},
            "end": {"dateTime": end_dt.isoformat(), "timeZone": tz},
            "attendees": [{"email": p} for p in participants if _EMAIL_RE.match(p)],
        }
        event = service.events().insert(
            calendarId=settings.GOOGLE_CALENDAR_ID, body=body
        ).execute()
        return {
            "status": "created",
            "event_id": event.get("id"),
            "html_link": event.get("htmlLink"),
        }
    except ImportError:
        return {"status": "simulated", "reason": "google-api-python-client not installed"}
    except Exception as exc:
        log.warning("calendar event creation failed: %s", exc)
        return {"status": "error", "reason": str(exc)[:200]}


def _send_meeting_email(title, start_dt, participants, user, calendar) -> dict:
    """Send a confirmation email (console backend in dev). Never raises."""
    recipients = [user.email] if user.email else []
    recipients += [p for p in participants if _EMAIL_RE.match(p) and p not in recipients]
    if not recipients:
        return {"status": "skipped", "reason": "no email recipients"}
    when = start_dt.strftime("%A, %d %B %Y at %H:%M") if start_dt else "a time to be confirmed"
    lines = [
        f"Hi {user.full_name or user.email},",
        "",
        f"Your meeting '{title}' has been scheduled for {when}.",
    ]
    if participants:
        lines.append(f"Participants: {', '.join(participants)}")
    if calendar.get("html_link"):
        lines.append(f"Calendar link: {calendar['html_link']}")
    lines += ["", "Sent automatically by the OpenVault assistant."]
    try:
        sent = send_mail(
            subject=f"Meeting scheduled: {title}",
            message="\n".join(lines),
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=recipients,
            fail_silently=False,
        )
        return {"status": "sent" if sent else "skipped", "recipients": recipients}
    except Exception as exc:
        log.warning("meeting email failed: %s", exc)
        return {"status": "error", "reason": str(exc)[:200], "recipients": recipients}


# ──────────────────────────── leave_request ─────────────────────────────────

def _leave_request(data: dict, user) -> dict:
    reason = (data.get("leave_reason") or data.get("reason") or "").strip()
    leave_date = (data.get("date") or "").strip()
    company_name = user.company.name if getattr(user, "company", None) else ""

    rel_path, filename = _generate_leave_docx(user, company_name, reason, leave_date)
    download_url = f"/api/workflow/forms/{user.id}/{filename}"
    if settings.PUBLIC_BASE_URL:
        download_url = settings.PUBLIC_BASE_URL + download_url

    return {
        "action": "leave_request",
        "status": "ok",
        "message": "Your leave application form is ready to download.",
        "filename": filename,
        "download_url": download_url,
        "leave_reason": reason,
        "leave_date": leave_date,
        # path is internal only; handy for logs/tests, harmless to the client.
        "stored_path": rel_path,
    }


def _generate_leave_docx(user, company_name, reason, leave_date) -> tuple[str, str]:
    """Build the auto-filled leave application .docx and return (relative_path, filename)."""
    from docx import Document  # python-docx (already a project dependency)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    today = timezone.localdate().strftime("%d %B %Y")
    full_name = user.full_name or user.email
    title = user.title or "Employee"

    doc = Document()
    heading = doc.add_heading("Leave Application", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run(f"Date: {today}").bold = True
    if company_name:
        doc.add_paragraph(f"Company: {company_name}")

    doc.add_paragraph("")
    doc.add_paragraph("To,")
    doc.add_paragraph("The Manager / HR Department")
    if company_name:
        doc.add_paragraph(company_name)

    doc.add_paragraph("")
    doc.add_paragraph("Subject: Application for Leave").runs[0].bold = True
    doc.add_paragraph("")

    body = doc.add_paragraph()
    body.add_run(f"Dear Sir/Madam,\n\n")
    sentence = f"I, {full_name} ({title}), would like to request leave"
    if leave_date:
        sentence += f" on {leave_date}"
    sentence += "."
    if reason:
        sentence += f" The reason for my leave is: {reason}."
    sentence += (
        " I would be grateful if you could approve my request. I will ensure that my"
        " responsibilities are managed appropriately during my absence."
    )
    body.add_run(sentence)

    doc.add_paragraph("")
    doc.add_paragraph("Thank you for your consideration.")
    doc.add_paragraph("")
    doc.add_paragraph("Sincerely,")
    name_p = doc.add_paragraph()
    name_p.add_run(full_name).bold = True
    doc.add_paragraph(title)

    for p in doc.paragraphs:
        for run in p.runs:
            if not run.font.size:
                run.font.size = Pt(11)

    user_dir = os.path.join(settings.MEDIA_ROOT, "forms", str(user.id))
    os.makedirs(user_dir, exist_ok=True)
    safe = _SAFE_NAME_RE.sub("_", (full_name or "employee").strip()) or "employee"
    filename = f"{uuid.uuid4().hex}_Leave_Application_{safe}.docx"
    abs_path = os.path.join(user_dir, filename)
    doc.save(abs_path)

    rel_path = os.path.join("forms", str(user.id), filename)
    return rel_path, filename
